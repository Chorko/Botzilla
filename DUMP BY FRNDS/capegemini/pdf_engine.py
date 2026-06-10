import whisperx
import torch
import json
from groq import Groq
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def get_chunks(segments, max_words= 1500):
    arr= []
    new_= ""
    res= 0
    for sg in segments:
        if "speaker" in sg:
            line= f"{sg[ 'start' ]:.3f} {sg[ 'speaker' ]} {sg[ 'text' ]}\n"
            count= len(line.split())
            if res+ count > max_words:
                arr.append(new_)
                new_= line
                res= count
            else:
                new_+= line
                res+= count
    if new_:
        arr.append(new_)
    return arr

def process_chunk(client, chunk_text):
    sys= "You are an AI meeting agent processing a segment of a transcript. Tasks: 1. Keep the exact speaker labels (e.g., SPEAKER_00). Do NOT change them to real names. 2. Fix broken pauses by the SAME speaker into one paragraph. 3. CRITICAL: Every single time a DIFFERENT speaker talks, you MUST start a new line. Output ONLY valid JSON matching this schema: { \"clean_transcript\": \"String\", \"action_items\": [ {\"task\": \"String\", \"owner\": \"String\", \"due_date\": \"String\"} ] }"
    new_= client.chat.completions.create(
        messages=[
            {"role": "system", "content": sys},
            {"role": "user", "content": chunk_text}
        ],
        model="llama-3.3-70b-versatile",
        response_format={"type": "json_object"}
    )
    return json.loads(new_.choices[ 0 ].message.content)

def gen_master(client, text):
    sys= "You are an executive analyzer. Tasks: 1. Read the full transcript and figure out the real names for SPEAKER_00, etc. 2. Generate summary. Output ONLY valid JSON matching this schema: { \"speaker_mapping\": {\"SPEAKER_00\": \"Real Name\"}, \"meeting_title\": \"String\", \"date\": \"YYYY-MM-DD\", \"attendees\": [\"String\"], \"metrics\": { \"duration_hours\": \"Number\", \"attendees_count\": \"Number\", \"decisions_count\": \"Number\", \"actions_count\": \"Number\" }, \"executive_summary\": \"String\", \"key_decisions\": [\"Point\"], \"key_discussion_points\": [\"Point\"], \"pending_issues\": [\"Point\"] }"
    new_= client.chat.completions.create(
        messages=[
            {"role": "system", "content": sys},
            {"role": "user", "content": text}
        ],
        model="llama-3.3-70b-versatile",
        response_format={"type": "json_object"}
    )
    return json.loads(new_.choices[ 0 ].message.content)

def process_full_audio(fp, hf_token, groq_key):
    print("[STATUS] transcribing_diarizing 10 Loading WhisperX models...")
    dev= "cuda" if torch.cuda.is_available() else "cpu"
    mdl= whisperx.load_model("tiny", dev)
    aud= whisperx.load_audio(fp)
    print("[STATUS] transcribing_diarizing 30 Transcribing vocal segments...")
    res= mdl.transcribe(aud, batch_size=16)
    print("[STATUS] transcribing_diarizing 50 Aligning phonetic timings...")
    aln, mtd= whisperx.load_align_model(language_code=res[ "language" ], device=dev)
    res= whisperx.align(res[ "segments" ], aln, mtd, aud, dev, return_char_alignments=False)
    print("[STATUS] transcribing_diarizing 70 Running Pyannote diarization...")
    dz_mdl= whisperx.diarize.DiarizationPipeline(token=hf_token, device=dev)
    dz_df= dz_mdl(aud)
    out= whisperx.assign_word_speakers(dz_df, res)
    arr= get_chunks(out[ "segments" ])
    client= Groq(api_key=groq_key)
    new_= ""
    final_actions= []
    print("[STATUS] resolving_names 80 Generating AI meeting summary...")
    for c in arr:
        data= process_chunk(client, c)
        new_+= data.get("clean_transcript", "") + "\n\n"
        final_actions.extend(data.get("action_items", []))
    master= gen_master(client, new_)
    map= master.get("speaker_mapping", {})
    for key, val in map.items():
        new_= new_.replace(key, val)
    print("[STATUS] finalizing 95 Compiling final Word document...")
    return {
        "summary": master,
        "clean_transcript": new_.strip(),
        "action_items": final_actions
    }

def gen_docx(data):
    new_= Document()
    style= new_.styles[ 'Normal' ]
    font= style.font
    font.name= 'Helvetica'
    font.size= Pt(11)
    head= new_.add_heading(data[ "summary" ][ "meeting_title" ], level= 1)
    head.alignment= WD_ALIGN_PARAGRAPH.CENTER
    for run in head.runs:
        run.font.color.rgb= RGBColor(0, 0, 0)
        run.font.name= 'Helvetica'
    new_.add_paragraph(f"Date: {data[ 'summary' ][ 'date' ]}")
    arr= ", ".join(data[ "summary" ][ "attendees" ])
    new_.add_paragraph(f"Attendees: {arr}")
    h2= new_.add_heading('Key Metrics', level= 2)
    for run in h2.runs:
        run.font.color.rgb= RGBColor(0, 0, 0)
    m= data[ "summary" ][ "metrics" ]
    new_.add_paragraph(f"Duration: {m[ 'duration_hours' ]} Hours | Attendees: {m[ 'attendees_count' ]} | Decisions: {m[ 'decisions_count' ]} | Actions: {m[ 'actions_count' ]}")
    h3= new_.add_heading('Executive Summary', level= 2)
    for run in h3.runs:
        run.font.color.rgb= RGBColor(0, 0, 0)
    new_.add_paragraph(data[ "summary" ][ "executive_summary" ])
    h4= new_.add_heading('Key Decisions', level= 2)
    for run in h4.runs:
        run.font.color.rgb= RGBColor(0, 0, 0)
    for d in data[ "summary" ][ "key_decisions" ]:
        new_.add_paragraph(f"- {d}")
    h5= new_.add_heading('Action Items', level= 2)
    for run in h5.runs:
        run.font.color.rgb= RGBColor(0, 0, 0)
    table= new_.add_table(rows= 1, cols= 3)
    table.style= 'Table Grid'
    hdr= table.rows[ 0 ].cells
    hdr[ 0 ].text= "Task"
    hdr[ 1 ].text= "Owner"
    hdr[ 2 ].text= "Deadline"
    for act in data[ "action_items" ]:
        res= table.add_row().cells
        res[ 0 ].text= act[ "task" ][ :45 ]
        res[ 1 ].text= act[ "owner" ][ :20 ]
        res[ 2 ].text= act[ "due_date" ][ :15 ]
    h6= new_.add_heading('Key Discussion Points', level= 2)
    for run in h6.runs:
        run.font.color.rgb= RGBColor(0, 0, 0)
    for p in data[ "summary" ][ "key_discussion_points" ]:
        new_.add_paragraph(f"- {p}")
    h7= new_.add_heading('Pending Issues', level= 2)
    for run in h7.runs:
        run.font.color.rgb= RGBColor(0, 0, 0)
    for p in data[ "summary" ][ "pending_issues" ]:
        new_.add_paragraph(f"- {p}")
    foot= new_.sections[ 0 ].footer
    foot_p= foot.paragraphs[ 0 ]
    foot_p.text= "Generated by SAKSHI LEDGER | Confidential"
    foot_p.alignment= WD_ALIGN_PARAGRAPH.CENTER
    out= "SAKSHI_LEDGER_Report.docx"
    new_.save(out)
    return out