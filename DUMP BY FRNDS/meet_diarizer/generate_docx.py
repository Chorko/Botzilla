import json
import sys
import os
import re
import io

# Force UTF-8 encoding on Windows consoles to prevent UnicodeEncodeError
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
sys.stderr = io.TextIOWrapper(sys.stderr.buffer, encoding='utf-8', errors='replace')
from collections import Counter, defaultdict
from docx import Document
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

def fmt_time(seconds: float) -> str:
    m = int(seconds // 60)
    s = int(seconds % 60)
    return f"{m:02d}:{s:02d}"

def fmt_duration(seconds: float) -> str:
    if seconds >= 3600:
        h = int(seconds // 3600)
        m = int((seconds % 3600) // 60)
        return f"{h}h {m}m"
    m = int(seconds // 60)
    s = int(seconds % 60)
    return f"{m}m {s}s"

def set_cell_shading(cell, color_hex):
    shading = cell._element.get_or_add_tcPr()
    shading_elem = shading.makeelement(qn('w:shd'), {
        qn('w:val'): 'clear',
        qn('w:color'): 'auto',
        qn('w:fill'): color_hex,
    })
    shading.append(shading_elem)

def set_col_widths(table, widths):
    for row in table.rows:
        for idx, width in enumerate(widths):
            row.cells[idx].width = width

def analyze_transcript(transcript: list) -> dict:
    """Analyze transcript to extract statistics using resolved speaker names."""
    speaker_stats = defaultdict(lambda: {"segments": 0, "words": 0, "talk_time": 0.0, "texts": []})
    total_duration = 0.0
    all_text = []

    for seg in transcript:
        spk = seg.get("speaker", "UNKNOWN")
        start = seg.get("start", 0.0)
        end = seg.get("end", 0.0)
        text = seg.get("text", "")
        duration = end - start

        speaker_stats[spk]["segments"] += 1
        speaker_stats[spk]["words"] += len(text.split())
        speaker_stats[spk]["talk_time"] += duration
        speaker_stats[spk]["texts"].append(text)
        all_text.append(text)

        if end > total_duration:
            total_duration = end

    return {
        "speaker_stats": dict(speaker_stats),
        "total_duration": total_duration,
    }

def add_heading_styled(doc, text, level, space_before=12, space_after=6):
    h = doc.add_heading(text, level=level)
    h.paragraph_format.space_before = Pt(space_before)
    h.paragraph_format.space_after = Pt(space_after)
    h.paragraph_format.keep_with_next = True
    
    # Configure colors based on level
    color = RGBColor(0x1A, 0x1A, 0x2E) # Indigo for H1
    if level == 2:
        color = RGBColor(0x44, 0x44, 0x88) # Muted blue for H2
    elif level == 3:
        color = RGBColor(0x66, 0x66, 0x99) # Slate gray for H3
        
    for run in h.runs:
        run.font.name = 'Arial'
        run.font.color.rgb = color
    return h

def generate_docx(json_path: str, output_path: str):
    with open(json_path, "r", encoding="utf-8") as f:
        data = json.load(f)

    metadata = data["metadata"]
    slides = data.get("slides", [])
    transcript = data.get("transcript", [])
    summary = data.get("summary", {})
    speaker_mapping = data.get("speaker_mapping", {})
    
    analysis = analyze_transcript(transcript)

    doc = Document()

    # -- Page margins (Standard Professional) --
    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)

    # -- Default font styling --
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Arial'
    font.size = Pt(10.5)
    font.color.rgb = RGBColor(0x33, 0x33, 0x33) # Charcoal body text

    # Color Palette Constants
    PRIMARY_COLOR = RGBColor(0x1A, 0x1A, 0x2E) # Indigo
    SECONDARY_COLOR = RGBColor(0x44, 0x44, 0x88) # Muted Blue
    ACCENT_COLOR = RGBColor(0x66, 0x66, 0x99) # Slate Gray
    
    # Build speaker color map for distinct styling in legend and transcript
    speakers = sorted(list(analysis["speaker_stats"].keys()))
    speaker_colors = [
        RGBColor(0x1B, 0x5E, 0x20),  # dark green
        RGBColor(0x0D, 0x47, 0xA1),  # dark blue
        RGBColor(0xBF, 0x36, 0x0C),  # dark orange
        RGBColor(0x4A, 0x14, 0x8C),  # dark purple
        RGBColor(0x00, 0x69, 0x5C),  # teal
        RGBColor(0x88, 0x00, 0x00),  # dark red
    ]
    speaker_color_map = {}
    for i, spk in enumerate(speakers):
        speaker_color_map[spk] = speaker_colors[i % len(speaker_colors)]

    # ================================================================
    # TITLE PAGE
    # ================================================================
    doc.add_paragraph("")
    doc.add_paragraph("")
    doc.add_paragraph("")
    
    title_text = summary.get("title", "Meeting Report")
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_p.add_run(title_text)
    title_run.font.size = Pt(26)
    title_run.font.bold = True
    title_run.font.name = 'Arial'
    title_run.font.color.rgb = PRIMARY_COLOR
    
    doc.add_paragraph("")
    
    subtitle_p = doc.add_paragraph()
    subtitle_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_run = subtitle_p.add_run("Automated AI Meeting Summarization & Slides Report")
    sub_run.font.size = Pt(13)
    sub_run.font.italic = True
    sub_run.font.color.rgb = SECONDARY_COLOR
    
    doc.add_paragraph("")
    doc.add_paragraph("")
    
    # -- Metadata table --
    meta_table = doc.add_table(rows=0, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_table.style = 'Light Shading Accent 1'
    
    duration_val = fmt_duration(analysis["total_duration"]) if analysis["total_duration"] > 0 else "N/A"
    meta_items = [
        ("Pipeline Mode", metadata.get("pipeline_mode", "N/A")),
        ("Source File", os.path.basename(metadata.get("video_file") or metadata.get("audio_file") or "N/A")),
        ("Primary Language", metadata.get("language", "N/A").upper()),
        ("Meeting Duration", duration_val),
        ("Total Participants", str(len(speakers))),
        ("Visual Slides Captured", str(metadata.get("num_slides_extracted", 0))),
    ]
    
    for label, value in meta_items:
        row = meta_table.add_row()
        row.cells[0].text = label
        row.cells[1].text = value
        row.cells[0].paragraphs[0].runs[0].font.bold = True
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(3)
                p.paragraph_format.space_after = Pt(3)
                for run in p.runs:
                    run.font.size = Pt(9.5)
                    run.font.color.rgb = RGBColor(0x44, 0x44, 0x44)
                    
    set_col_widths(meta_table, [Inches(2.5), Inches(3.5)])
    
    doc.add_page_break()

    # ================================================================
    # EXECUTIVE SUMMARY & DECISIONS
    # ================================================================
    add_heading_styled(doc, "Executive Summary", level=1)
    
    exec_text = summary.get("executive_summary", "")
    if exec_text:
        p = doc.add_paragraph()
        r = p.add_run(exec_text)
        r.font.size = Pt(10.5)
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.space_after = Pt(12)
        
    # Key Decisions
    decisions = summary.get("key_decisions", [])
    if decisions:
        add_heading_styled(doc, "Key Decisions Made", level=2)
        for dec in decisions:
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_after = Pt(4)
            r = p.add_run(dec)
            r.font.size = Pt(10)
            
    doc.add_paragraph("")

    # ================================================================
    # ACTION ITEMS
    # ================================================================
    action_items = summary.get("action_items", [])
    if action_items:
        add_heading_styled(doc, "Action Items & Assignments", level=1)
        
        # Create Table
        action_table = doc.add_table(rows=1, cols=3)
        action_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        action_table.style = 'Light Shading Accent 1'
        
        # Format Headers
        hdr_cells = action_table.rows[0].cells
        headers = ["Action Task Description", "Assignee", "Priority"]
        for idx, h_text in enumerate(headers):
            hdr_cells[idx].text = h_text
            set_cell_shading(hdr_cells[idx], "1A1A2E") # Indigo Background
            p = hdr_cells[idx].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if idx > 0 else WD_ALIGN_PARAGRAPH.LEFT
            for r in p.runs:
                r.font.bold = True
                r.font.size = Pt(9.5)
                r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF) # White text
                
        # Populate Items
        for row_idx, item in enumerate(action_items):
            row = action_table.add_row()
            cells = row.cells
            cells[0].text = item.get("task", "")
            cells[1].text = item.get("assignee", "Unassigned")
            
            prio = item.get("priority", "Medium").capitalize()
            cells[2].text = prio
            
            # Format and Shading
            shading_color = "FFFFFF" if row_idx % 2 == 0 else "F2F4F7" # Alternating rows
            for c_idx, cell in enumerate(cells):
                set_cell_shading(cell, shading_color)
                p = cell.paragraphs[0]
                p.paragraph_format.space_before = Pt(4)
                p.paragraph_format.space_after = Pt(4)
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if c_idx > 0 else WD_ALIGN_PARAGRAPH.LEFT
                
                for r in p.runs:
                    r.font.size = Pt(9.5)
                    # Priority color highlights
                    if c_idx == 2:
                        r.font.bold = True
                        if prio == "High":
                            r.font.color.rgb = RGBColor(0xB7, 0x1C, 0x1C) # Dark Red
                        elif prio == "Medium":
                            r.font.color.rgb = RGBColor(0xE6, 0x51, 0x00) # Orange
                        else:
                            r.font.color.rgb = RGBColor(0x55, 0x55, 0x55) # Gray
                            
        set_col_widths(action_table, [Inches(3.8), Inches(1.4), Inches(1.0)])
        doc.add_paragraph("")

    # ================================================================
    # KEY DISCUSSION TOPICS
    # ================================================================
    topics = summary.get("topics", [])
    if topics:
        add_heading_styled(doc, "Key Discussion Topics", level=1)
        for top in topics:
            add_heading_styled(doc, top.get("name", "Topic"), level=2)
            p = doc.add_paragraph()
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.space_after = Pt(8)
            r = p.add_run(top.get("summary", ""))
            r.font.size = Pt(10)
            
        doc.add_page_break()

    # ================================================================
    # SPEAKER PARTICIPATION
    # ================================================================
    add_heading_styled(doc, "Speaker Participation Summary", level=1)
    
    spk_table = doc.add_table(rows=1, cols=4)
    spk_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    spk_table.style = 'Light Shading Accent 1'
    
    # Headers
    hdr = spk_table.rows[0].cells
    headers = ["Speaker Name", "Talk Duration", "Participation %", "Words Spoken"]
    for idx, text in enumerate(headers):
        hdr[idx].text = text
        set_cell_shading(hdr[idx], "444488") # Muted Blue Header
        p = hdr[idx].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if idx > 0 else WD_ALIGN_PARAGRAPH.LEFT
        for r in p.runs:
            r.font.bold = True
            r.font.size = Pt(9.5)
            r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            
    total_talk_time = sum(s["talk_time"] for s in analysis["speaker_stats"].values())
    
    for row_idx, spk in enumerate(speakers):
        stats = analysis["speaker_stats"][spk]
        pct = (stats["talk_time"] / total_talk_time * 100) if total_talk_time > 0 else 0
        
        row = spk_table.add_row()
        cells = row.cells
        cells[0].text = spk
        cells[1].text = fmt_duration(stats["talk_time"])
        cells[2].text = f"{pct:.1f}%"
        cells[3].text = f"{stats['words']:,}"
        
        shading_color = "FFFFFF" if row_idx % 2 == 0 else "F2F4F7"
        for c_idx, cell in enumerate(cells):
            set_cell_shading(cell, shading_color)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if c_idx > 0 else WD_ALIGN_PARAGRAPH.LEFT
            
            for r in p.runs:
                r.font.size = Pt(9.5)
                if c_idx == 0:
                    r.font.bold = True
                    r.font.color.rgb = speaker_color_map.get(spk, PRIMARY_COLOR)
                    
    set_col_widths(spk_table, [Inches(2.5), Inches(1.2), Inches(1.3), Inches(1.2)])
    doc.add_page_break()

    # ================================================================
    # FULL MEETING TRANSCRIPT WITH INLINE SLIDES
    # ================================================================
    add_heading_styled(doc, "Full Meeting Transcript", level=1)
    
    # Speaker Legend
    legend_p = doc.add_paragraph()
    legend_p.paragraph_format.space_after = Pt(12)
    legend_run = legend_p.add_run("Speaker Legend:  ")
    legend_run.font.bold = True
    legend_run.font.size = Pt(9)
    for spk in speakers:
        spk_run = legend_p.add_run(f"  ● {spk}  ")
        spk_run.font.bold = True
        spk_run.font.size = Pt(9)
        spk_run.font.color.rgb = speaker_color_map.get(spk, PRIMARY_COLOR)
        
    doc.add_paragraph("")
    
    last_slide_embedded = None
    
    for seg in transcript:
        spk = seg.get("speaker", "UNKNOWN")
        start = seg.get("start", 0.0)
        end = seg.get("end", 0.0)
        text = seg.get("text", "")
        seg_slides = seg.get("slides_shown", [])
        
        # Segment paragraph
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.15
        
        # Speaker tag
        spk_run = p.add_run(f"{spk}")
        spk_run.font.bold = True
        spk_run.font.size = Pt(10)
        spk_run.font.color.rgb = speaker_color_map.get(spk, PRIMARY_COLOR)
        
        # Timestamp
        ts_run = p.add_run(f"  [{fmt_time(start)} - {fmt_time(end)}]")
        ts_run.font.size = Pt(8.5)
        ts_run.font.color.rgb = ACCENT_COLOR
        
        # Text content
        text_run = p.add_run(f"\n{text}")
        text_run.font.size = Pt(10)
        
        # Inline slide rendering (embedded directly below the speaker turn where it was active)
        if seg_slides:
            for slide_path in seg_slides:
                # Avoid embedding the exact same slide consecutively in short succession
                if slide_path != last_slide_embedded and os.path.isfile(slide_path):
                    # Add caption
                    caption_p = doc.add_paragraph()
                    caption_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    caption_p.paragraph_format.space_before = Pt(8)
                    caption_p.paragraph_format.space_after = Pt(2)
                    
                    # Clean filename timestamp extraction for display
                    base_slide_name = os.path.basename(slide_path)
                    ts_match = re.search(r"slide_(\d+m\d+s)", base_slide_name)
                    ts_label = ts_match.group(1).replace('m', ':').replace('s', '') if ts_match else "Active"
                    
                    cap_run = caption_p.add_run(f"🖼️ [Slide Active at {ts_label}]")
                    cap_run.font.size = Pt(8.5)
                    cap_run.font.bold = True
                    cap_run.font.italic = True
                    cap_run.font.color.rgb = SECONDARY_COLOR
                    
                    # Embed Image
                    try:
                        doc.add_picture(slide_path, width=Inches(4.8))
                        img_p = doc.paragraphs[-1]
                        img_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        img_p.paragraph_format.space_after = Pt(12)
                        last_slide_embedded = slide_path
                    except Exception as e:
                        print(f"⚠️ Failed to insert image {slide_path}: {e}")

    # ================================================================
    # FOOTER
    # ================================================================
    doc.add_paragraph("")
    footer_p = doc.add_paragraph()
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = footer_p.add_run("— End of Meeting Summary Report —")
    r.font.color.rgb = ACCENT_COLOR
    r.font.size = Pt(9.5)
    r.font.italic = True

    doc.save(output_path)
    print(f"✅ Word document generated successfully: {output_path}")

if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("Usage: python generate_docx.py <botzilla_output.json> [output.docx]")
        sys.exit(1)

    json_path = sys.argv[1]
    if len(sys.argv) >= 3:
        out_path = sys.argv[2]
    else:
        out_path = json_path.replace("_botzilla.json", "_report.docx").replace(".json", ".docx")

    generate_docx(json_path, out_path)
