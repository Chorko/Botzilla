import sys
import os
import whisperx
import torch
import json
import time
from groq import Groq
from google import genai
from google.genai import types
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def update_p(start_val, end_val):
    for i in range(start_val, end_val+ 1):
        print(f"[STATUS] progression {i}%", flush=True)
        time.sleep(0.02)

def get_chunks(segments, max_words=1500):
    arr= []
    new_= ""
    res= 0
    for sg in segments:
        if "speaker" in sg:
            line= f"[{sg[ 'start' ]:.2f} - {sg[ 'end' ]:.2f}] {sg[ 'speaker' ]} : {sg[ 'text' ]}\n"
            count= len(line.split())
            if res+ count > max_words:
                arr.append(new_)
                new_= line
                res= count
            else:
                new_+= line
                res+= count
    if new_:
        arr.append(new_)
    return arr

def process_chunk(client, chunk_text):
    sys_= "You are an AI meeting agent processing a segment of a transcript.\nTasks:\n1. Keep the exact speaker labels (e.g., SPEAKER_00) and exact timestamps. Do NOT change them.\n2. Fix broken pauses by the SAME speaker into one paragraph.\n3. CRITICAL: Every single time a DIFFERENT speaker talks, you MUST start a new line. NEVER merge two different speakers into the same block of text.\n4. Extract action items.\nOutput ONLY valid JSON matching this schema:\n{\n  \"clean_transcript\": \"String\",\n  \"action_items\": [\n    {\"task\": \"String\", \"owner\": \"String\", \"due_date\": \"String\"}\n  ]\n}"
    new_= client.chat.completions.create(
        messages=[
            {"role": "system", "content": sys_},
            {"role": "user", "content": chunk_text}
        ],
        model="llama-3.3-70b-versatile",
        response_format={"type": "json_object"}
    )
    return json.loads(new_.choices[ 0 ].message.content)

def gen_master(client, gemini_key, text, routing_flag):
    sys_= "You are an executive analyzer.\nTasks:\n1. Read the full transcript, analyze the context deeply, and deduce the real names for SPEAKER_00, SPEAKER_01, etc.\n2. Look at how people greet each other and respond to figure out who is who.\n3. Generate a formal architecture sync summary.\n4. Append exact spoken quotes in brackets for all key points.\n5. CRITICAL: NEVER use double quotes inside your text values. You MUST use single quotes for all internal quotes to prevent JSON crashes.\nOutput ONLY valid JSON matching this schema:\n{\n  \"speaker_mapping\": {\"SPEAKER_00\": \"Real Name\", \"SPEAKER_01\": \"Real Name\"},\n  \"meeting_title\": \"String\",\n  \"date\": \"YYYY-MM-DD\",\n  \"attendees\": [\"String\"],\n  \"metrics\": {\n    \"duration_hours\": \"Number\",\n    \"attendees_count\": \"Number\",\n    \"decisions_count\": \"Number\",\n    \"actions_count\": \"Number\"\n  },\n  \"executive_summary\": \"String\",\n  \"key_decisions\": [\"Point 1. ['Exact quote']\"],\n  \"key_discussion_points\": [\"Point 1. ['Exact quote']\"],\n  \"pending_issues\": [\"Point 1. ['Exact quote']\"]\n}"
    if routing_flag == "gemini":
        client_g= genai.Client(api_key=gemini_key)
        new_= client_g.models.generate_content(
            model='gemini-2.5-flash',
            contents=text,
            config=types.GenerateContentConfig(
                system_instruction=sys_,
                response_mime_type="application/json",
            ),
        )
        return json.loads(new_.text)
    else:
        new_= client.chat.completions.create(
            messages=[
                {"role": "system", "content": sys_},
                {"role": "user", "content": text}
            ],
            model="llama-3.3-70b-versatile",
            response_format={"type": "json_object"}
        )
        return json.loads(new_.choices[ 0 ].message.content)

def process_full_audio(fp, hf_token, groq_key, gemini_key, task_dir):
    update_p(1, 10)
    print("[STATUS] Loaded WhisperX Frame Models", flush=True)
    dev= "cuda" if torch.cuda.is_available() else "cpu"
    mdl= whisperx.load_model("tiny", dev)
    aud= whisperx.load_audio(fp)
    
    update_p(11, 30)
    print("[STATUS] Transcribing audio frequencies...", flush=True)
    res= mdl.transcribe(aud, batch_size=16)
    
    client= Groq(api_key=groq_key)
    
    update_p(31, 40)
    raw_= " ".join([sg[ "text" ] for sg in res[ "segments" ][ :40 ]])
    sys_= "You are an AI data analyst processing a meeting transcript. Determine the total number of distinct speakers. Normal calls have 2 to 5 speakers. NEVER output a number higher than 8. Output ONLY a single integer digit. Do not write any words."
    
    new_= client.chat.completions.create(
        messages=[
            {"role": "system", "content": sys_},
            {"role": "user", "content": raw_}
        ],
        model="llama-3.3-70b-versatile"
    )
    
    val_= new_.choices[ 0 ].message.content.strip()
    count_= 3
    if val_.isdigit():
        count_= int(val_)
        
    print(f"[STATUS] Detected total of {count_} people speaking in this session", flush=True)
    
    routing_flag= "llama"
    if count_ > 4:
        routing_flag= "gemini"
        print("[STATUS] Scale threat detected: Loaded Gemini Core Engine for processing", flush=True)
    else:
        print("[STATUS] Loaded Llama Core Engine for processing", flush=True)
        
    update_p(41, 55)
    aln, mtd= whisperx.load_align_model(language_code=res[ "language" ], device=dev)
    res= whisperx.align(res[ "segments" ], aln, mtd, aud, dev, return_char_alignments=False)
    
    update_p(56, 75)
    dz_mdl= whisperx.diarize.DiarizationPipeline(token=hf_token, device=dev)
    dz_df= dz_mdl(aud, min_speakers=1, max_speakers=8)
    out= whisperx.assign_word_speakers(dz_df, res)
    
    arr= []
    new_= None
    for sg in out[ "segments" ]:
        if "speaker" in sg:
            if new_ == None:
                new_= {
                    "speaker": sg[ "speaker" ],
                    "text": sg[ "text" ].strip(),
                    "start": sg[ "start" ],
                    "end": sg[ "end" ]
                }
            elif new_[ "speaker" ] == sg[ "speaker" ]:
                new_[ "text" ]+= " " + sg[ "text" ].strip()
                new_[ "end" ]= sg[ "end" ]
            else:
                arr.append(new_)
                new_= {
                    "speaker": sg[ "speaker" ],
                    "text": sg[ "text" ].strip(),
                    "start": sg[ "start" ],
                    "end": sg[ "end" ]
                }
    if new_ != None:
        arr.append(new_)
        
    out[ "segments" ]= arr
    
    print("\n--- [RAW DIARIZATION CHECK] ---", flush=True)
    for sg in out[ "segments" ]:
        print(f"[{sg[ 'start' ]:.2f} - {sg[ 'end' ]:.2f}] {sg[ 'speaker' ]}: {sg[ 'text' ]}", flush=True)
    print("-------------------------------\n", flush=True)
    
    res= get_chunks(out[ "segments" ])
    
    new_= ""
    final_actions= []
    
    update_p(76, 85)
    for c in res:
        data= process_chunk(client, c)
        new_+= data.get("clean_transcript", "") + "\n\n"
        final_actions.extend(data.get("action_items", []))
            
    master= gen_master(client, gemini_key, new_, routing_flag)
    map_= master.get("speaker_mapping", {})
    for key, val in map_.items():
        new_= new_.replace(key, val)
        
    update_p(86, 95)
    ui_seg= []
    for sg in out[ "segments" ]:
        ui_seg.append({
            "speaker": map_.get(sg[ "speaker" ], sg[ "speaker" ]),
            "text": sg[ "text" ],
            "startTime": sg[ "start" ],
            "endTime": sg[ "end" ]
        })
        
    ui_dec= []
    for d in master.get("key_decisions", []):
        ui_dec.append({ "decision": d, "owner": "Team" })
        
    ui_act= []
    for a in final_actions:
        ui_act.append({
            "task": a.get("task", "Task"),
            "owner": a.get("owner", "Unassigned"),
            "priority": "High",
            "deadline": a.get("due_date", "TBD")
        })
        
    ui_kp= []
    for p in master.get("key_discussion_points", []):
        ui_kp.append({ "category": "Discussion", "point": p })

    ui_json= {
        "title": master.get("meeting_title", "Botzilla Sync"),
        "date": master.get("date", "Today"),
        "executiveSummary": master.get("executive_summary", ""),
        "speakerMapping": map_,
        "segments": ui_seg,
        "decisions": ui_dec,
        "actionItems": ui_act,
        "keyPoints": ui_kp,
        "metrics": master.get("metrics", {})
    }
    
    with open(os.path.join(task_dir, "report.json"), "w") as f:
        json.dump(ui_json, f)
        
    update_p(96, 100)
    return {
        "summary": master,
        "clean_transcript": new_.strip(),
        "action_items": final_actions
    }

def gen_docx(data, task_dir):
    new_= Document()
    style= new_.styles[ 'Normal' ]
    font= style.font
    font.name= 'Helvetica'
    font.size= Pt(11)
    
    head= new_.add_heading(data[ "summary" ].get("meeting_title", "Botzilla Sync"), level= 1)
    head.alignment= WD_ALIGN_PARAGRAPH.CENTER
    for run in head.runs:
        run.font.color.rgb= RGBColor(30, 58, 138)
        run.font.name= 'Helvetica'
        
    new_.add_paragraph(f"Date: {data[ 'summary' ].get('date', 'Today')}")
    arr= ", ".join(data[ "summary" ].get("attendees", []))
    new_.add_paragraph(f"Attendees: {arr}")
    
    h2= new_.add_heading('Key Metrics', level= 2)
    for run in h2.runs:
        run.font.color.rgb= RGBColor(37, 99, 235)
    m= data[ "summary" ].get("metrics", {})
    new_.add_paragraph(f"Duration: {m.get('duration_hours', 0)} Hours | Attendees: {m.get('attendees_count', 0)} | Decisions: {m.get('decisions_count', 0)} | Actions: {m.get('actions_count', 0)}")
    
    h3= new_.add_heading('Executive Summary', level= 2)
    for run in h3.runs:
        run.font.color.rgb= RGBColor(37, 99, 235)
    new_.add_paragraph(data[ "summary" ].get("executive_summary", "N/A"))
    
    h4= new_.add_heading('Key Decisions', level= 2)
    for run in h4.runs:
        run.font.color.rgb= RGBColor(37, 99, 235)
    for d in data[ "summary" ].get("key_decisions", []):
        new_.add_paragraph(f"- {d}")
        
    h5= new_.add_heading('Action Items', level= 2)
    for run in h5.runs:
        run.font.color.rgb= RGBColor(37, 99, 235)
        
    table= new_.add_table(rows= 1, cols= 3)
    table.style= 'Table Grid'
    hdr= table.rows[ 0 ].cells
    hdr[ 0 ].text= "Task"
    hdr[ 1 ].text= "Owner"
    hdr[ 2 ].text= "Deadline"
    
    for act in data[ "action_items" ]:
        res= table.add_row().cells
        res[ 0 ].text= act.get("task", "")[ :45 ]
        res[ 1 ].text= act.get("owner", "")[ :20 ]
        res[ 2 ].text= act.get("due_date", "")[ :15 ]
        
    h6= new_.add_heading('Key Discussion Points', level= 2)
    for run in h6.runs:
        run.font.color.rgb= RGBColor(37, 99, 235)
    for p in data[ "summary" ].get("key_discussion_points", []):
        new_.add_paragraph(f"- {p}")
        
    h7= new_.add_heading('Pending Issues', level= 2)
    for run in h7.runs:
        run.font.color.rgb= RGBColor(37, 99, 235)
    for p in data[ "summary" ].get("pending_issues", []):
        new_.add_paragraph(f"- {p}")
        
    foot= new_.sections[ 0 ].footer
    foot_p= foot.paragraphs[ 0 ]
    foot_p.text= "Generated by Botzilla | Confidential"
    foot_p.alignment= WD_ALIGN_PARAGRAPH.CENTER
    
    out= os.path.join(task_dir, "Botzilla_Report.docx")
    new_.save(out)
    return out

if __name__ == "__main__":
    audio_path= sys.argv[ 1 ]
    hf_token= sys.argv[ 2 ]
    groq_key= sys.argv[ 3 ]
    gemini_key= sys.argv[ 4 ]
    task_dir= sys.argv[ 5 ]
    
    new_= process_full_audio(audio_path, hf_token, groq_key, gemini_key, task_dir)
    output_docx= gen_docx(new_, task_dir)
    print(f"[SUCCESS] {output_docx}", flush=True)